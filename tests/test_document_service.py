import pytest
from fastapi import HTTPException

from nyayaai.services.document_service import (
    clean_text,
    extract_text,
    split_clauses,
    validate_upload,
)


def test_clean_text_normalizes_whitespace() -> None:
    assert clean_text(" hello   world\n\n\n next ") == "hello world\n\n next"


def test_split_clauses_tracks_pages_and_clause_numbers() -> None:
    text = "[PAGE 1]\n1. Payment\nThe customer pays monthly.\n2. Termination\nThirty days notice."
    clauses = split_clauses(text)
    assert [clause.id for clause in clauses] == ["C1", "C2"]
    assert clauses[0].page == 1
    assert clauses[1].text.startswith("2. Termination")


def test_validate_upload_rejects_unknown_extension() -> None:
    with pytest.raises(HTTPException) as exc_info:
        validate_upload("malware.exe", "application/octet-stream", b"data")
    assert exc_info.value.status_code == 400


def test_validate_upload_rejects_empty_file() -> None:
    with pytest.raises(HTTPException, match="empty"):
        validate_upload("agreement.txt", "text/plain", b"")


def test_validate_upload_rejects_invalid_pdf_signature() -> None:
    with pytest.raises(HTTPException, match="invalid"):
        validate_upload("agreement.pdf", "application/pdf", b"not-a-pdf")


def test_extract_text_reads_utf8_text() -> None:
    assert extract_text("agreement.txt", b"Payment is due monthly.") == "Payment is due monthly."


def test_extract_text_rejects_non_utf8_text() -> None:
    with pytest.raises(HTTPException, match="UTF-8"):
        extract_text("agreement.txt", b"\xff\xfe")


def test_validate_upload_accepts_valid_docx() -> None:
    from io import BytesIO

    from docx import Document

    buffer = BytesIO()
    document = Document()
    document.add_paragraph("Payment is due monthly.")
    document.save(buffer)
    assert validate_upload(
        "agreement.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        buffer.getvalue(),
    ) == ".docx"


def test_extract_text_reads_docx() -> None:
    from io import BytesIO

    from docx import Document

    buffer = BytesIO()
    document = Document()
    document.add_paragraph("Payment is due monthly.")
    document.save(buffer)
    assert "Payment is due monthly." in extract_text("agreement.docx", buffer.getvalue())


def test_split_clauses_handles_plain_text_without_headings() -> None:
    clauses = split_clauses("A simple legal paragraph without numbered headings.")
    assert len(clauses) == 1
    assert clauses[0].id == "C1"
